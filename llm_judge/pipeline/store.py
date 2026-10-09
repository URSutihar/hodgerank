"""Dual-format, crash-safe, Excel-lock-tolerant result store.

Three files, each with a distinct job:

  judgments.jsonl     append-only log. THE authoritative record. One line per
                      judgement, fsync'd, so a kill -9 loses at most the line
                      being written.
  judgments.parquet   columnar snapshot for fast downstream analysis.
  judgments.xlsx      human-readable snapshot, formatted, with a Summary sheet.

Why the Excel file is a *snapshot* and not the source of truth: on Windows,
Excel takes an exclusive lock while a workbook is open, so any write fails with
PermissionError. Rather than crash or lose data, we keep appending to the JSONL,
mark Excel dirty, and rewrite it on the next flush that succeeds. Close Excel and
the very next flush catches it up to the exact judgement count reached. Nothing
is ever lost, and the run never stops for a locked spreadsheet.
"""

from __future__ import annotations

import json
import os
import time
from pathlib import Path
from typing import Any, Iterable

import pandas as pd

from .config import Config

# Column order for the human-readable sheet.
COLUMNS = [
    "row", "pair_id", "presentation", "left_image", "right_image",
    "left_age", "right_age", "age_gap", "true_older_side",
    "model_side", "model_choice", "model_correct", "model_confidence",
    "human_side", "human_choice", "human_correct", "annotator",
    "model_agrees_with_human",
    "model_name", "quantization", "prompt_version",
    "latency_s", "raw_response", "error", "timestamp_iso",
]

EXCEL_MIN_INTERVAL_S = 5.0   # never rewrite the workbook more often than this


class JudgmentStore:
    def __init__(self, cfg: Config):
        self.cfg = cfg
        cfg.ensure_dirs()
        self.jsonl_path = cfg.jsonl_path
        self.parquet_path = cfg.parquet_path
        self.excel_path = cfg.excel_path

        self._records: list[dict[str, Any]] = []
        self._fh = None
        self._excel_dirty = False
        self._last_excel_write = 0.0
        self._excel_blocked_since: float | None = None

        self._load_existing()

    # ── startup / resume ────────────────────────────────────────────────
    def _load_existing(self) -> None:
        if not self.jsonl_path.exists():
            return
        good, bad = [], 0
        with open(self.jsonl_path, "r", encoding="utf-8") as fh:
            for line in fh:
                line = line.strip()
                if not line:
                    continue
                try:
                    good.append(json.loads(line))
                except json.JSONDecodeError:
                    bad += 1          # torn final line from a hard kill
        self._records = good
        if bad:
            # A hard kill can leave a partially written final line. Rewrite the
            # log from the records that parsed, so the damage does not persist.
            tmp = self.jsonl_path.with_suffix(".jsonl.tmp")
            with open(tmp, "w", encoding="utf-8") as out:
                for rec in good:
                    out.write(json.dumps(rec, ensure_ascii=False) + "\n")
            os.replace(tmp, self.jsonl_path)
        if good:
            print(f"resuming: {len(good):,} judgements already in "
                  f"{self.jsonl_path.name}"
                  + (f" ({bad} torn line repaired)" if bad else ""))

    @property
    def done_keys(self) -> set[str]:
        """Keys already judged, so a resumed run skips them."""
        return {f"{r.get('pair_id')}::{r.get('presentation')}" for r in self._records}

    def __len__(self) -> int:
        return len(self._records)

    # ── writing ─────────────────────────────────────────────────────────
    def _open(self):
        if self._fh is None:
            self._fh = open(self.jsonl_path, "a", encoding="utf-8")
        return self._fh

    def append(self, record: dict[str, Any]) -> None:
        """Record one judgement. Durable before this returns."""
        record.setdefault("row", len(self._records) + 1)
        self._records.append(record)
        fh = self._open()
        fh.write(json.dumps(record, ensure_ascii=False) + "\n")
        fh.flush()
        os.fsync(fh.fileno())          # survive power loss, not just process death
        self._excel_dirty = True

    def to_frame(self) -> pd.DataFrame:
        if not self._records:
            return pd.DataFrame(columns=COLUMNS)
        df = pd.DataFrame(self._records)
        for c in COLUMNS:
            if c not in df.columns:
                df[c] = pd.NA
        extra = [c for c in df.columns if c not in COLUMNS]
        return df[COLUMNS + extra]

    # ── snapshots ───────────────────────────────────────────────────────
    def flush(self, force: bool = False) -> None:
        """Refresh the parquet and Excel snapshots from the JSONL record set."""
        if not self._records:
            return
        df = self.to_frame()
        self._write_parquet(df)
        if force or (time.time() - self._last_excel_write) >= EXCEL_MIN_INTERVAL_S:
            self._write_excel(df)

    def _write_parquet(self, df: pd.DataFrame) -> bool:
        tmp = self.parquet_path.with_suffix(".parquet.tmp")
        try:
            df.to_parquet(tmp, index=False)
            os.replace(tmp, self.parquet_path)
            return True
        except Exception as exc:                       # noqa: BLE001
            print(f"  [warn] parquet write failed: {exc}")
            tmp.unlink(missing_ok=True)
            return False

    def _write_excel(self, df: pd.DataFrame) -> bool:
        """Atomic rewrite. Tolerates the workbook being open in Excel.

        Returns True if the file on disk is now current.
        """
        tmp = self.excel_path.with_name(self.excel_path.stem + ".tmp.xlsx")
        try:
            with pd.ExcelWriter(tmp, engine="openpyxl") as xl:
                df.to_excel(xl, sheet_name="judgments", index=False)
                self._summary_frame(df).to_excel(
                    xl, sheet_name="summary", index=False)
                _style(xl, df)
            os.replace(tmp, self.excel_path)          # atomic where the OS allows
            self._last_excel_write = time.time()
            self._excel_dirty = False
            if self._excel_blocked_since is not None:
                waited = time.time() - self._excel_blocked_since
                print(f"  [ok] Excel unlocked — synced {len(df):,} rows "
                      f"(was blocked {waited:.0f}s)")
                self._excel_blocked_since = None
            return True
        except PermissionError:
            # Workbook is open (typically Excel on Windows). Data is safe in the
            # JSONL; we simply try again next flush.
            tmp.unlink(missing_ok=True)
            self._excel_dirty = True
            if self._excel_blocked_since is None:
                self._excel_blocked_since = time.time()
                print(f"  [info] {self.excel_path.name} is open — run continues, "
                      f"will sync when you close it")
            return False
        except Exception as exc:                       # noqa: BLE001
            tmp.unlink(missing_ok=True)
            self._excel_dirty = True
            print(f"  [warn] Excel write failed: {exc}")
            return False

    def _summary_frame(self, df: pd.DataFrame) -> pd.DataFrame:
        scored = df[df["model_correct"].notna()]
        hscored = df[df["human_correct"].notna()]
        both = df[df["model_side"].notna() & df["human_side"].notna()]
        errs = int(df["error"].notna().sum()) if "error" in df else 0

        rows: list[tuple[str, Any]] = [
            ("generated", pd.Timestamp.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")),
            ("model", df["model_name"].dropna().iloc[-1] if df["model_name"].notna().any() else ""),
            ("quantization", df["quantization"].dropna().iloc[-1] if df["quantization"].notna().any() else ""),
            ("judgements recorded", len(df)),
            ("errors", errs),
            ("model accuracy vs true age",
             f"{100*scored['model_correct'].mean():.2f}%" if len(scored) else "—"),
            ("human accuracy vs true age",
             f"{100*hscored['human_correct'].mean():.2f}%" if len(hscored) else "—"),
            ("model agrees with human",
             f"{100*both['model_agrees_with_human'].mean():.2f}%"
             if len(both) and both["model_agrees_with_human"].notna().any() else "—"),
            ("model left-side rate",
             f"{100*(df['model_side'] == 'left').mean():.2f}%" if len(df) else "—"),
            ("median latency (s)",
             f"{df['latency_s'].median():.2f}" if df["latency_s"].notna().any() else "—"),
        ]
        # accuracy broken out by difficulty
        if len(scored):
            for lo, hi, label in [(1, 3, "1-3"), (4, 7, "4-7"), (8, 15, "8-15"),
                                  (16, 30, "16-30"), (31, 200, "31+")]:
                sub = scored[(scored["age_gap"] >= lo) & (scored["age_gap"] <= hi)]
                if len(sub):
                    rows.append((f"  model acc · gap {label} yr",
                                 f"{100*sub['model_correct'].mean():.2f}% (n={len(sub)})"))
        return pd.DataFrame(rows, columns=["metric", "value"])

    # ── shutdown ────────────────────────────────────────────────────────
    def close(self) -> None:
        """Final sync. Retries a locked workbook so the file ends up current."""
        if self._fh is not None:
            self._fh.flush()
            os.fsync(self._fh.fileno())
            self._fh.close()
            self._fh = None
        if not self._records:
            return

        df = self.to_frame()
        self._write_parquet(df)

        for attempt in range(3):
            self._write_excel(df)
            if not self._excel_dirty:
                break
            if attempt == 0:
                print(f"  waiting for {self.excel_path.name} to be closed "
                      f"(retrying ~10s)…")
            time.sleep(5)

        if self._excel_dirty:
            print(f"\n  ! {self.excel_path.name} is still open, so it is behind by "
                  f"some rows.\n    Nothing is lost — all {len(df):,} judgements are in "
                  f"{self.jsonl_path.name}.\n    Close the workbook and run:  "
                  f"python cli.py sync")
        else:
            print(f"  synced {len(df):,} rows to {self.excel_path.name}")


def _style(xl, df: pd.DataFrame) -> None:
    """Freeze panes, autofilter, sane column widths, colour the correctness flags."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = xl.sheets["judgments"]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "pair_id": 46, "left_image": 34, "right_image": 34, "human_choice": 34,
        "model_choice": 34, "raw_response": 40, "timestamp_iso": 22,
        "model_name": 20, "error": 28, "annotator": 12, "presentation": 13,
    }
    for i, col in enumerate(df.columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 13)
        ws.cell(row=1, column=i).font = Font(bold=True)
        ws.cell(row=1, column=i).alignment = Alignment(horizontal="center")

    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    for name in ("model_correct", "human_correct"):
        if name not in df.columns:
            continue
        ci = list(df.columns).index(name) + 1
        for ri, val in enumerate(df[name], start=2):
            if pd.isna(val):
                continue
            ws.cell(row=ri, column=ci).fill = green if float(val) == 1.0 else red

    s = xl.sheets["summary"]
    s.column_dimensions["A"].width = 34
    s.column_dimensions["B"].width = 30
    for c in ("A1", "B1"):
        s[c].font = Font(bold=True)


def rebuild_snapshots(cfg: Config) -> int:
    """Regenerate parquet + Excel from the JSONL. Used by `cli.py sync`."""
    store = JudgmentStore(cfg)
    if not len(store):
        print("no judgements recorded yet")
        return 0
    df = store.to_frame()
    pq_ok = store._write_parquet(df)
    xl_ok = store._write_excel(df)
    if xl_ok and pq_ok:
        print(f"rebuilt {cfg.excel_path.name} and {cfg.parquet_path.name} "
              f"from {len(df):,} judgements")
    else:
        if not xl_ok:
            print(f"could NOT write {cfg.excel_path.name} — close it in Excel "
                  f"and run `sync` again")
        if not pq_ok:
            print(f"could NOT write {cfg.parquet_path.name}")
        print(f"all {len(df):,} judgements remain safe in {cfg.jsonl_path.name}")
    return len(df)
